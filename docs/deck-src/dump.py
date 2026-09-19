"""Dump every text shape's geometry and run properties for the overflow check."""
import json, sys
from pptx import Presentation
from pptx.util import Emu

prs = Presentation(sys.argv[1])
out = []
for i, sl in enumerate(prs.slides, 1):
    shapes = []
    for sh in sl.shapes:
        d = {"kind": sh.shape_type.__str__() if sh.shape_type is not None else "?",
             "x": Emu(sh.left).inches, "y": Emu(sh.top).inches,
             "w": Emu(sh.width).inches, "h": Emu(sh.height).inches}
        if sh.has_text_frame and sh.text_frame.text.strip():
            tf = sh.text_frame
            d["text"] = tf.text
            runs = [r for p in tf.paragraphs for r in p.runs]
            if runs:
                r = runs[0]
                d["size"] = r.font.size.pt if r.font.size else 18
                d["face"] = r.font.name or "Trebuchet MS"
                d["bold"] = bool(r.font.bold)
            sp = [p.line_spacing for p in tf.paragraphs if p.line_spacing]
            if sp:
                d["lineSpacing"] = sp[0].pt if hasattr(sp[0], "pt") else sp[0]
        shapes.append(d)
    out.append({"n": i, "shapes": shapes})
json.dump(out, open(sys.argv[2], "w"))
print("dumped", len(out), "slides")
